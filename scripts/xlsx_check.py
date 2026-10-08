"""openpyxl helpers for the round-trip test (no wOS code involved, so it is an independent reader).
  python scripts/xlsx_check.py make <out.xlsx>             writes a workbook with formulas and prints the expected values as JSON
  python scripts/xlsx_check.py read <file.xlsx>            prints every cell's cached value and formula as JSON
"""
import json, sys, datetime
import openpyxl

def make(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Deals'
    rows = [('Deal', 'Stage', 'Amount', 'Probability', 'Close'),
            ('Staff plan', 'Won', 18000, 1.0, datetime.date(2026, 10, 16)),
            ('Family plan', 'Proposal', 12000, 0.5, datetime.date(2026, 10, 21)),
            ('Screening day', 'Qualified', 6200, 0.25, datetime.date(2026, 11, 12)),
            ('Driver checkups', 'Proposal', 22400, 0.5, datetime.date(2026, 11, 3))]
    for r in rows: ws.append(r)
    exp = {}
    for i in range(2, 6):
        ws[f'F{i}'] = f'=C{i}*D{i}'
        exp[f'Deals!F{i}'] = rows[i - 1][2] * rows[i - 1][3]
        ws[f'G{i}'] = f'=IF(B{i}="Won","Closed","Open")'
        exp[f'Deals!G{i}'] = 'Closed' if rows[i - 1][1] == 'Won' else 'Open'
    ws['C6'] = '=SUM(C2:C5)'; exp['Deals!C6'] = sum(r[2] for r in rows[1:])
    ws['F6'] = '=SUM(F2:F5)'; exp['Deals!F6'] = sum(r[2] * r[3] for r in rows[1:])
    ws['H2'] = '=ROUND(AVERAGE(C2:C5),2)'; exp['Deals!H2'] = round(sum(r[2] for r in rows[1:]) / 4, 2)
    ws['H3'] = '=MAX(E2:E5)-MIN(E2:E5)'; exp['Deals!H3'] = (max(r[4] for r in rows[1:]) - min(r[4] for r in rows[1:])).days
    ws['C2'].number_format = '"$"#,##0'
    ws.freeze_panes = 'A2'
    s = wb.create_sheet('Summary')
    s['A1'] = 'Proposal total'; s['B1'] = '=SUMIF(Deals!B2:B5,"Proposal",Deals!C2:C5)'; exp['Summary!B1'] = 34400
    s['A2'] = 'Deals'; s['B2'] = '=COUNTA(Deals!A2:A5)'; exp['Summary!B2'] = 4
    s['A3'] = 'Lookup'; s['B3'] = '=VLOOKUP("Family plan",Deals!A2:C5,3,FALSE)'; exp['Summary!B3'] = 12000
    s['A4'] = 'Weighted share'; s['B4'] = '=ROUND(Deals!F6/Deals!C6,4)'; exp['Summary!B4'] = round(exp['Deals!F6'] / exp['Deals!C6'], 4)
    s['A5'] = 'Label'; s['B5'] = '=CONCATENATE("Q4: ",TEXT(Deals!C6,"#,##0"))'; exp['Summary!B5'] = 'Q4: 58,600'
    s['A6'] = 'Months'; s['B6'] = '=DATEDIF(DATE(2026,1,1),Deals!E3,"m")'; exp['Summary!B6'] = 9
    wb.save(path)
    print(json.dumps(exp))

def read(path):
    out = {}
    vals = openpyxl.load_workbook(path, data_only=True)
    forms = openpyxl.load_workbook(path, data_only=False)
    for ws in vals.worksheets:
        fws = forms[ws.title]
        for row in ws.iter_rows():
            for c in row:
                f = fws[c.coordinate].value
                v = c.value
                if isinstance(v, (datetime.datetime, datetime.date)): v = v.isoformat()[:10]
                if v is None and f is None: continue
                out[f'{ws.title}!{c.coordinate}'] = {'value': v, 'formula': f if isinstance(f, str) and f.startswith('=') else None}
    print(json.dumps(out))

{'make': make, 'read': read}[sys.argv[1]](sys.argv[2])
